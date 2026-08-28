"""生成"云雀科技有限公司"脱敏样例库（内容虚构、结构等价、可复现）。

覆盖：
- 多格式: MD / DOCX / PPTX / PDF（文本型）
- 多结构: 长文档（制度手册）、短文档（通知/说明）、标题层级、表格、PPT 备注页
- 权限: department + confidentiality 元数据（用于 MetadataFilters 联调）

用法: uv run python -m rag.scripts.generate_sample
产物: sample_data/ 目录 + manifest.csv
"""
from __future__ import annotations

import csv
import shutil
from pathlib import Path

from docx import Document as DocxDocument
from pptx import Presentation
from pptx.util import Inches, Pt
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas

PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent.parent
ROOT = PROJECT_ROOT / "sample_data"

# ---------------------------------------------------------------------------
# 内容（全部虚构）
# ---------------------------------------------------------------------------

POLICY_MD = """# 云雀科技有限公司制度手册（2026 版）

> 本手册适用于云雀科技有限公司全体员工，自 2026 年 1 月 1 日起施行。
> 如需调整或解释，请咨询人力资源部或行政部。

## 第一章 总则

第一条 为规范公司管理、保障员工权益，依据国家相关法律法规制定本手册。

第二条 本手册适用于公司全体员工，包括正式员工、试用期员工与劳务派遣人员。

## 第二章 考勤管理

### 2.1 工作时间

| 时段 | 时间 |
| --- | --- |
| 上午 | 09:00 – 12:00 |
| 午休 | 12:00 – 13:30 |
| 下午 | 13:30 – 18:30 |

弹性工作制：员工可在 08:30 – 09:30 之间到岗，对应 18:00 – 19:00 之间离岗，
每日在岗时长不少于 8 小时（含午休 1.5 小时中的固定工时部分按部门规定执行）。

### 2.2 迟到与早退

单月累计迟到 3 次以内不作处理；超过 3 次，每次扣减当月绩效 0.5 分。

### 2.3 请假流程

- 事假、病假、年假均须提前在 OA 系统提交申请，注明事由与时长。
- 请假 1 天以内由直属主管审批；1 至 3 天由部门负责人审批；3 天以上需分管 VP 审批。
- 病假超过 3 天的，须提供三甲医院出具的诊断证明。

### 2.4 加班管理

- 因工作需要加班，须**提前一天**在 OA 系统提交加班申请，经直属主管审批后生效；
  未经审批的自行加班不计入加班工时。
- 工作日加班按 1.5 倍工资、休息日加班按 2 倍工资、法定节假日加班按 3 倍工资计算。

## 第三章 差旅管理

### 3.1 差旅申请

出差前须在 OA 提交《出差申请单》，注明目的地、事由、起止日期与预算，
经部门负责人审批后方可出行。

### 3.2 住宿标准

| 城市类别 | 住宿标准（元/晚） |
| --- | --- |
| 一线城市（北上广深） | 不超过 500 |
| 其他城市 | 不超过 350 |
| 港澳台及境外 | 不超过 800 |

超出部分自理，特殊情况须事先报分管 VP 特批。

### 3.3 交通标准

- 高铁：二等座/一等座均可报销；行程超过 6 小时的，可报销商务座。
- 飞机：单程 4 小时以上的航线方可乘飞机，经济舱报销。
- 市内交通：凭发票实报实销，单日上限 200 元。

## 第四章 费用报销

### 4.1 报销流程

1. 在 OA 系统填写《费用报销单》，选择费用类型并关联出差申请/采购申请编号；
2. 上传发票影像与付款凭证；
3. 直属主管审批 → 财务部复核 → 出纳付款；
4. 正常付款周期为审批通过后 5 个工作日内。

### 4.2 发票要求

- 报销须提供**增值税发票**（普通发票或专用发票均可），发票抬头必须为
  **"云雀科技有限公司"**，并填写公司税号；抬头不符或缺失税号的发票不予报销。
- 电子发票须提供 PDF 原件与查验截图；纸质发票须保持票面完整、无涂改。
- 单张发票金额超过 5000 元的，须同时提供采购合同或审批单。

## 第五章 信息安全

### 5.1 数据分级

公司数据按敏感程度分为三级：**公开（public）**、**内部（internal）**、**秘密（secret）**。

| 级别 | 示例 | 访问范围 |
| --- | --- | --- |
| 公开 | 公司官网内容、对外宣传材料 | 全员 + 外部 |
| 内部 | 制度手册、内部通知、部门资料 | 公司全员（或按部门） |
| 秘密 | 客户合同、薪酬数据、未公开财报 | 授权人员 |

### 5.2 终端安全

- 办公电脑须开启磁盘加密与锁屏（锁屏超时不超过 10 分钟）；
- 禁止在个人设备上存储秘密级数据；内部及以上数据拷贝须经信息安全部审批。

### 5.3 涉密信息处理

- 秘密级文件禁止外发、禁止通过私人邮箱/网盘传输，传输须使用公司加密通道；
- 纸质秘密文件废弃时须碎纸处理；
- 违反信息安全规定的，视情节给予警告、记过直至解除劳动合同，并保留追究法律责任的权利。

## 第六章 附则

本手册解释权归公司人力资源部与行政部。与法律法规冲突的条款以法律法规为准。
"""

ONBOARDING_DOCX = [
    ("heading", "新员工入职指南"),
    ("para", "欢迎加入云雀科技有限公司！本指南帮助你顺利度过入职第一天与试用期。"),
    ("heading2", "一、入职前准备"),
    ("para", "人力资源部会在你入职前 3 个工作日发送《入职确认函》，请携带身份证原件、学历学位证书复印件、银行卡信息（用于工资发放）办理入职。"),
    ("heading2", "二、入职第一天要做的事"),
    ("para", "1. 到前台领取工牌，办理门禁卡与人脸识别录入；"),
    ("para", "2. 到 IT 支持处领取办公电脑，开通公司邮箱、OA 账号、VPN 账号；"),
    ("para", "3. 参加 10:00 的新员工入职培训（公司介绍、制度要点、信息安全须知）；"),
    ("para", "4. 由直属主管介绍团队与岗位职责，明确试用期目标；"),
    ("para", "5. 在 OA 系统完成个人信息补充（紧急联系人、银行卡、报销账户）。"),
    ("heading2", "三、常用系统"),
    ("para", "OA 系统：请假、报销、加班、出差申请。"),
    ("para", "企业微信：日常沟通与会议。"),
    ("para", "代码仓库与文档平台：产品、研发、项目资料。"),
    ("heading2", "四、福利概览"),
    ("para", "五险一金按国家规定缴纳，另有补充商业保险、年度体检、弹性工作制、节日礼品与团建经费。"),
]

RND_FLOW_DOCX = [
    ("heading", "研发流程规范"),
    ("para", "本文档适用于云雀科技有限公司所有研发项目，是需求到上线的全过程基线。"),
    ("heading2", "1. 需求评审"),
    ("para", "需求进入研发前必须完成需求评审：产品经理讲解需求背景与验收标准，研发与测试评估可行性与工作量，评审结论写入需求单。未评审的需求不得排期。"),
    ("heading2", "2. 设计评审"),
    ("para", "涉及架构变更、数据库变更、对外接口变更的需求，必须提交设计文档并完成设计评审。设计评审至少包含 2 名高级工程师参加。"),
    ("heading2", "3. 代码评审"),
    ("para", "所有代码变更必须通过代码评审才能合入主干："),
    ("para", "- 普通模块：至少 1 名 reviewer 通过；"),
    ("para", "- 核心模块（支付、权限、数据存储）：至少 2 名 reviewer 通过；"),
    ("para", "- 评审前必须通过 CI 静态检查与单元测试，覆盖率不得低于 80%。"),
    ("heading2", "4. 发布流程"),
    ("para", "发布采用灰度发布：先在预发环境验证，再按 10% → 50% → 100% 逐步放量。每次发布必须有回滚预案，发布窗口为工作日 10:00–16:00。"),
    ("heading2", "5. 故障响应"),
    ("para", "线上故障按 P0/P1/P2 分级：P0（核心业务不可用）须在 15 分钟内拉起应急群，30 分钟内给出初步结论，并完成事故复盘报告。"),
]

PRODUCT_PLAN_SLIDES = [
    (
        "2026 年 Q1 产品规划",
        "云雀科技有限公司 产品部\n汇报人：产品部\n密级：内部",
        None,
    ),
    (
        "一、Q1 核心目标",
        "1. 新增用户 100 万，重点投入华东市场；\n2. 核心功能（云盘、文档协作）周留存提升至 45%；\n3. 商业化收入达成 800 万元。",
        "目标拆解：华东市场投放预算占总预算 40%。",
    ),
    (
        "二、重点投入",
        "华东市场品牌投放\n企业版销售团队扩编 30 人\n文档协作 2.0 版本研发",
        "文档协作 2.0 预计 3 月底上线。",
    ),
    (
        "三、OKR 表格",
        "| O | KR | 负责人 |\n|---|---|---|\n| O1 用户增长 | KR1 新增 100 万用户 | 增长组 |\n| O1 用户增长 | KR2 华东获客成本低于 35 元 | 增长组 |\n| O2 留存提升 | KR1 周留存 45% | 产品组 |\n| O3 商业化 | KR1 收入 800 万 | 商业化组 |",
        None,
    ),
    (
        "四、风险与依赖",
        "风险：华东市场竞争加剧、投放成本上升。\n依赖：数据平台口径对齐、客服团队扩容。",
        None,
    ),
]

CLOUD_INTRO_PDF = [
    ("云雀云服务产品介绍", "云雀科技有限公司 市场部  2026 年 2 月"),
    ("一、产品概述", "云雀云服务是面向中小企业的云计算平台，提供云服务器、对象存储、云数据库、文档协作与音视频会议等一站式服务。"),
    ("二、核心能力", "1. 弹性计算：分钟级开通，支持按量付费与包年包月；\n2. 对象存储：单文件最大 5TB，跨区域容灾；\n3. 云数据库：MySQL 与 PostgreSQL 兼容，自动备份与主备切换；\n4. 文档协作：多人实时编辑、权限细粒度控制、全程审计。"),
    ("三、服务可用性（SLA）", "云服务器、对象存储、云数据库的月度服务可用性承诺为 99.9%。未达标的，按服务等级协议提供补偿。"),
    ("四、安全与合规", "通过等保三级测评，支持私有网络（VPC）、密钥管理与全链路加密。数据存储于境内数据中心。"),
    ("五、联系我们", "官网：www.yunque.example.com  客服热线：400-000-0000（7×24 小时）"),
]

HOLIDAY_NOTICE_MD = """# 关于 2026 年五一假期安排的通知

各位同事：

根据国家法定节假日安排，结合公司实际，现将 2026 年五一劳动节放假安排通知如下：

## 一、放假时间

**5 月 1 日（周五）至 5 月 5 日（周二）放假调休，共 5 天。**

4 月 26 日（周日）、5 月 9 日（周六）正常上班（补班）。

## 二、注意事项

1. 放假前请关闭办公区域电源与门窗，重要文件锁入文件柜；
2. 值班安排见 OA 系统值班表，值班人员保持手机畅通；
3. 假期出行注意安全，如遇突发情况请及时联系行政部（内线 8001）。

祝大家节日愉快！

云雀科技有限公司 行政部
2026 年 4 月 10 日
"""

EXPENSE_NOTE_MD = """# 差旅报销补充说明（2026 年 3 月修订）

> 本说明是对《公司制度手册（2026 版）》第三章、第四章的补充，两者冲突时以本说明为准。

## 一、住宿发票要求

- 住宿发票抬头必须为"云雀科技有限公司"，并填写税号；
- 平台（携程、飞猪等）开具的发票若抬头不符，须同时提供行程单与支付记录方可报销。

## 二、打车与市内交通

- 22:00 之后的加班打车，凭企业微信加班打卡记录实报实销，不设单日上限；
- 普通市内交通仍按单日 200 元上限执行。

## 三、报销时限

- 出差归来后 **15 个自然日内**须提交报销单，逾期需书面说明原因并经财务总监审批。

## 四、常见驳回原因

1. 发票抬头不符或缺失税号；
2. 缺少出差申请单编号关联；
3. 超标住宿未附特批记录。

云雀科技有限公司 财务部
2026 年 3 月 15 日
"""

MANIFEST = [
    # path, doc_id, title, department, confidentiality, doc_type
    ("公司制度手册.md", "policy_handbook", "公司制度手册（2026版）", "行政部", "internal", "md"),
    ("新员工入职指南.docx", "onboarding_guide", "新员工入职指南", "人力资源部", "public", "docx"),
    ("研发流程规范.docx", "rnd_flow", "研发流程规范", "技术部", "internal", "docx"),
    ("2026年Q1产品规划.pptx", "product_plan_q1", "2026年Q1产品规划", "产品部", "internal", "pptx"),
    ("云雀云服务产品介绍.pdf", "cloud_intro", "云雀云服务产品介绍", "市场部", "public", "pdf"),
    ("关于2026年五一假期安排的通知.md", "holiday_notice", "五一假期安排通知", "行政部", "public", "md"),
    ("差旅报销补充说明.md", "expense_note", "差旅报销补充说明", "财务部", "internal", "md"),
]


# ---------------------------------------------------------------------------
# 生成器
# ---------------------------------------------------------------------------

def write_md(rel: str, content: str) -> None:
    p = ROOT / rel
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(content, encoding="utf-8")


def write_docx(rel: str, blocks: list[tuple[str, str]]) -> None:
    doc = DocxDocument()
    for kind, text in blocks:
        if kind == "heading":
            doc.add_heading(text, level=0)
        elif kind == "heading2":
            doc.add_heading(text, level=1)
        else:
            doc.add_paragraph(text)
    p = ROOT / rel
    p.parent.mkdir(parents=True, exist_ok=True)
    doc.save(p)


def write_pptx(rel: str, slides: list[tuple[str, str, str | None]]) -> None:
    prs = Presentation()
    for title, body, notes in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])  # title + content
        slide.shapes.title.text = title
        tf = slide.placeholders[1].text_frame
        first = True
        for line in body.split("\n"):
            para = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            para.text = line
            para.font.size = Pt(16)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    p = ROOT / rel
    p.parent.mkdir(parents=True, exist_ok=True)
    prs.save(p)


def write_pdf(rel: str, sections: list[tuple[str, str]]) -> None:
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    p = ROOT / rel
    p.parent.mkdir(parents=True, exist_ok=True)
    c = canvas.Canvas(str(p), pagesize=A4)
    width, height = A4
    y = height - 60
    for title, body in sections:
        c.setFont("STSong-Light", 16)
        c.drawString(50, y, title)
        y -= 28
        c.setFont("STSong-Light", 11)
        for line in body.split("\n"):
            if y < 60:
                c.showPage()
                y = height - 60
            c.drawString(50, y, line)
            y -= 18
        y -= 12
    c.save()


def write_manifest() -> None:
    p = ROOT / "manifest.csv"
    with open(p, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(["path", "doc_id", "title", "department", "confidentiality", "doc_type"])
        w.writerows(MANIFEST)


def main() -> None:
    if ROOT.exists():
        shutil.rmtree(ROOT)
    ROOT.mkdir(parents=True)

    write_md("公司制度手册.md", POLICY_MD)
    write_docx("新员工入职指南.docx", ONBOARDING_DOCX)
    write_docx("研发流程规范.docx", RND_FLOW_DOCX)
    write_pptx("2026年Q1产品规划.pptx", PRODUCT_PLAN_SLIDES)
    write_pdf("云雀云服务产品介绍.pdf", CLOUD_INTRO_PDF)
    write_md("关于2026年五一假期安排的通知.md", HOLIDAY_NOTICE_MD)
    write_md("差旅报销补充说明.md", EXPENSE_NOTE_MD)
    write_manifest()

    print(f"样例库已生成: {ROOT}")
    for rel, *_ in MANIFEST:
        print(f"  - {rel}")


if __name__ == "__main__":
    main()
